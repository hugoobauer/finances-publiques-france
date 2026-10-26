import openpyxl,glob,csv,sys,re,statistics as st
d=sys.argv[1]; out=sys.argv[2]
OECD={'Australia','Austria','Belgium (Flemish)','Belgium (French)','Canada','Chile','Colombia','Costa Rica','Czech Republic','Czechia','Denmark','Estonia','Finland','France','Germany','Greece','Hungary','Iceland','Ireland','Israel','Italy','Japan','Korea, Rep. of','Latvia','Lithuania','Luxembourg','Mexico','Netherlands','New Zealand','Norway','Poland','Portugal','Slovak Republic','Slovenia','Spain','Sweden','Switzerland','Türkiye','England','Northern Ireland','United States'}
EU={'Austria','Belgium (Flemish)','Belgium (French)','Bulgaria','Croatia','Cyprus','Czech Republic','Czechia','Denmark','Estonia','Finland','France','Germany','Greece','Hungary','Ireland','Italy','Latvia','Lithuania','Luxembourg','Malta','Netherlands','Poland','Portugal','Romania','Slovak Republic','Slovenia','Spain','Sweden'}
PANEL={'France','Germany','Italy','Spain','Netherlands','Belgium (Flemish)','Belgium (French)','Sweden','Denmark'}
rows_out=[]
for fn in sorted(glob.glob(d+'/*.xlsx')):
    ws=openpyxl.load_workbook(fn,data_only=True).worksheets[0]
    title=[c for c in list(ws.iter_rows(values_only=True))[1] if c][0]
    vals={}
    for r in ws.iter_rows(values_only=True):
        r=[c for c in r if c is not None]
        for i in range(len(r)-1):
            if isinstance(r[i],str) and isinstance(r[i+1],(int,float)) and r[i].strip() and not r[i].strip().startswith(('Exhibit','Note')):
                name=re.sub(r'\s*\(\d+\)$','',r[i].replace('⋈','').strip()).strip(); vals[name]=r[i+1]; break
    print(title, len(vals))
    for g,S in [('OECD',OECD),('EU',EU),('Panel',PANEL)]:
        v={c:x for c,x in vals.items() if c in S}
        if 'France' not in v: continue
        s=sorted(v.values(),reverse=True)
        rows_out.append([title,g,len(v),round(st.mean(v.values()),1),st.median(v.values()),v['France'],s.index(v['France'])+1,'; '.join(f'{k} {x}' for k,x in sorted(v.items(),key=lambda kv:-kv[1]))])
        print(' ',rows_out[-1][:7])
with open(out,'w',newline='') as f:
    w=csv.writer(f); w.writerow(['exhibit','group','n','mean_unweighted','median','france','fr_rank','countries_scores']); w.writerows(rows_out)
