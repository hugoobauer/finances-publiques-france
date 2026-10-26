"""PISA 2025 Volume I (Annex B1) extraction: means, trends, low/top performers, SES gradient.
Usage: python3 -I pisa.py <research_dir>"""
import openpyxl, csv, sys, statistics as st
R = sys.argv[1]; raw = R + '/data/raw/13/'
A = openpyxl.load_workbook(raw + 'pisa25_mrq53f.xlsx', read_only=True, data_only=True)
B = openpyxl.load_workbook(raw + 'pisa25_k68msa.xlsx', read_only=True, data_only=True)
OECD = ['Australia','Austria','Belgium','Canada','Chile','Colombia','Costa Rica','Czechia','Denmark','Estonia','Finland','France',
        'Germany','Greece','Hungary','Iceland','Ireland','Israel','Italy','Japan','Korea','Latvia','Lithuania','Luxembourg','Mexico',
        'Netherlands','New Zealand','Norway','Poland','Portugal','Slovak Republic','Slovenia','Spain','Sweden','Switzerland','Türkiye',
        'United Kingdom','United States']
EU = ['Austria','Belgium','Bulgaria','Croatia','Cyprus','Czechia','Denmark','Estonia','Finland','France','Germany','Greece','Hungary',
      'Ireland','Italy','Latvia','Lithuania','Luxembourg','Malta','Netherlands','Poland','Portugal','Romania','Slovak Republic',
      'Slovenia','Spain','Sweden']
PANEL = ['France','Germany','Italy','Spain','Netherlands','Belgium','Sweden','Denmark']
def rows(wb, name):
    out = {}
    for r in wb[name].iter_rows(values_only=True):
        if r[0] and isinstance(r[0], str):
            out[r[0].replace('*', '').strip()] = (r, r[0].endswith('*'))
    return out
def num(x):
    return x if isinstance(x, (int, float)) else None
data = {}  # (country, var) -> value
flag = {}
dom = {'science': ('1', '36', '33', '2'), 'reading': ('2', '37', '34', '3'), 'math': ('3', '38', '35', '4')}
for d, (tm, tt, tlp, tses) in dom.items():
    for c, (r, ast) in rows(A, 'Table I.B1.2a.' + tm).items():
        data[(c, d + '_2025')] = num(r[1]); flag[c] = ast or flag.get(c, False)
    # trend tables: columns PISA 2000/2003/2006... mean at odd positions
    hdr = None
    for i, r in enumerate(A['Table I.B1.2a.' + tt].iter_rows(values_only=True)):
        if i == 7: hdr = r
        if i >= 9 and r[0]:
            c = r[0].replace('*', '').strip()
            for j, h in enumerate(hdr):
                if h and isinstance(h, str) and h.strip().startswith('PISA') and '\n' not in h and j < len(r):
                    data[(c, d + '_' + h.strip()[5:9])] = num(r[j])
    # low / top performers: the 2025 pair sits after the cycles
    hdr6 = None
    for i, r in enumerate(A['Table I.B1.2a.' + tlp].iter_rows(values_only=True)):
        if i == 6: hdr6 = r
        if i >= 9 and r[0]:
            c = r[0].replace('*', '').strip()
            for j, h in enumerate(hdr6):
                if h and isinstance(h, str) and h.startswith('Proficiency level in PISA'):
                    y = h.strip()[-4:]
                    data[(c, d + '_low_' + y)] = num(r[j]); data[(c, d + '_top_' + y)] = num(r[j + 2])
    for i, r in enumerate(B['Table I.B1.2b.' + tses].iter_rows(values_only=True)):
        if i >= 9 and r[0]:
            c = r[0].replace('*', '').strip()
            data[(c, d + '_ses_r2')] = num(r[1]); data[(c, d + '_ses_slope')] = num(r[4])
            data[(c, d + '_ses_gap_q4q1')] = num(r[19])
countries = sorted(set(k[0] for k in data))
vars_ = sorted(set(k[1] for k in data))
with open(R + '/data/13_pisa2025_country_values.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['country', 'oecd', 'eu27', 'panel', 'asterisk_flag'] + vars_)
    for c in countries:
        if c.startswith(('Information', 'Note', 'Notes', '1.', '2.', 'Source')): continue
        w.writerow([c, c in OECD, c in EU, c in PANEL, flag.get(c, '')] + [('' if data.get((c, v)) is None else round(data[(c, v)], 2)) for v in vars_])
def grp(v, members, higher_better=True):
    vals = {c: data.get((c, v)) for c in members}; vals = {c: x for c, x in vals.items() if x is not None}
    if 'France' not in vals: return None
    s = sorted(vals.values(), reverse=higher_better)
    return len(vals), st.mean(vals.values()), st.median(vals.values()), vals['France'], s.index(vals['France']) + 1
out = []
for d in dom:
    for v, hb in [(d + '_2025', True), (d + '_2022', True), (d + '_2015', True), (d + '_2006', True),
                  (d + '_low_2025', False), (d + '_top_2025', True), (d + '_low_2015', False), (d + '_top_2015', True),
                  (d + '_ses_r2', False), (d + '_ses_slope', False), (d + '_ses_gap_q4q1', False)]:
        for g, m in [('OECD38', OECD), ('EU27', EU), ('Panel8', PANEL)]:
            s = grp(v, m, hb)
            if s:
                out.append([v, g, s[0], round(s[1], 1), round(s[2], 1), round(s[3], 1), s[4], 'higher=better' if hb else 'lower=better (rank 1 = lowest)'])
oa = {}
for name, d in [('Table I.B1.2a.36', 'science'), ('Table I.B1.2a.37', 'reading'), ('Table I.B1.2a.38', 'math')]:
    pass
with open(R + '/data/13_pisa2025_group_stats.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['variable', 'group', 'n', 'mean_unweighted', 'median', 'france', 'fr_rank', 'rank_convention']); w.writerows(out)
for r in out: print(r)
print()
for c in PANEL + ['OECD average', 'OECD average-23', 'OECD average-35']:
    print(c, {v: (round(data[(c, v)], 1) if data.get((c, v)) is not None else None) for v in
              ['math_2025', 'reading_2025', 'science_2025', 'math_2015', 'reading_2015', 'science_2015', 'science_2006', 'math_2006', 'reading_2006', 'math_2003', 'reading_2000',
               'math_low_2025', 'math_top_2025', 'reading_low_2025', 'reading_top_2025', 'science_low_2025', 'science_top_2025', 'math_ses_r2', 'math_ses_slope']})
print('asterisk:', [c for c in countries if flag.get(c)])
